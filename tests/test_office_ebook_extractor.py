# seam-scope: implementation-infrastructure (public format extractor fixtures)
from __future__ import annotations

import zipfile
from pathlib import Path

from docops.extractors import ExtractorPolicy
from docops.extractors.office_ebook import OfficeEbookExtractor
from docops.ir import validate_ir_document


def _policy() -> ExtractorPolicy:
    return ExtractorPolicy(
        rights_ref="rights-office",
        source_id="source-office",
        source_revision_id="revision-office",
        required_fidelity="structured-native",
    )


def test_docx_preserves_headings_paragraphs_and_tables(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "guide.docx"
    document = Document()
    document.add_heading("Guide", level=1)
    document.add_paragraph("Use a token.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "mode"
    table.cell(1, 1).text = "strict"
    document.save(source)

    result = OfficeEbookExtractor().extract(source, _policy(), {})

    assert result.document is not None
    assert [block.kind for block in result.document.blocks][:2] == ["heading", "paragraph"]
    assert any(block.kind == "table" and "strict" in str(block.structured) for block in result.document.blocks)
    assert validate_ir_document(result.document).ok


def test_epub_follows_spine_and_preserves_chapters(tmp_path: Path) -> None:
    source = tmp_path / "book.epub"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr(
            "META-INF/container.xml",
            "<container><rootfiles><rootfile full-path='OEBPS/content.opf'/></rootfiles></container>",
        )
        archive.writestr(
            "OEBPS/content.opf",
            "<package xmlns='http://www.idpf.org/2007/opf'><manifest>"
            "<item id='c1' href='chapter1.xhtml' media-type='application/xhtml+xml'/>"
            "<item id='c2' href='chapter2.xhtml' media-type='application/xhtml+xml'/>"
            "</manifest><spine><itemref idref='c1'/><itemref idref='c2'/></spine></package>",
        )
        archive.writestr("OEBPS/chapter1.xhtml", "<html><body><h1>One</h1><p>First.</p></body></html>")
        archive.writestr("OEBPS/chapter2.xhtml", "<html><body><h1>Two</h1><p>Second.</p></body></html>")

    result = OfficeEbookExtractor().extract(source, _policy(), {})

    assert result.document is not None
    assert [block.text for block in result.document.blocks if block.kind == "heading"] == ["One", "Two"]
    assert [block.locators[0]["chapter"] for block in result.document.blocks if block.kind == "heading"] == [1, 2]


def test_archive_path_traversal_is_rejected_before_materialization(tmp_path: Path) -> None:
    source = tmp_path / "hostile.epub"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("../escape.txt", "must not extract")

    result = OfficeEbookExtractor().extract(source, _policy(), {})

    assert result.document is None
    assert result.receipt.status == "failed"
    assert result.receipt.errors[0]["code"] == "invalid_archive"
    assert not (tmp_path / "escape.txt").exists()
